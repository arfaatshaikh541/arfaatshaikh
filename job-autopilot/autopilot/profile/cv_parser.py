"""Deterministic CV text extraction and structuring.

Everything produced here has status EXTRACTED: the candidate must confirm
(or correct) it before it is used to answer an application question.
Nothing is invented: every extracted fact carries the CV excerpt it came from.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

SUPPORTED_MIME = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "txt",
}


class CVParseError(Exception):
    pass


def detect_kind(filename: str, data: bytes) -> str:
    fn = filename.lower()
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:2] == b"PK" and fn.endswith(".docx"):
        return "docx"
    if fn.endswith(".txt"):
        try:
            data.decode("utf-8")
            return "txt"
        except UnicodeDecodeError:
            pass
    raise CVParseError("Unsupported CV format. Upload a PDF, DOCX or UTF-8 TXT file.")


def extract_text(filename: str, data: bytes) -> str:
    kind = detect_kind(filename, data)
    if kind == "pdf":
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((p.extract_text() or "") for p in reader.pages)
        except Exception as e:
            raise CVParseError(f"Could not read PDF: {e}") from e
        if not text.strip():
            raise CVParseError("PDF contains no extractable text (scanned image?). OCR is NOT SUPPORTED.")
        return text
    if kind == "docx":
        import docx

        try:
            d = docx.Document(io.BytesIO(data))
        except Exception as e:
            raise CVParseError(f"Could not read DOCX: {e}") from e
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    return data.decode("utf-8")


@dataclass
class ExtractedFact:
    category: str
    key: str
    value: str
    excerpt: str
    group_key: str | None = None


@dataclass
class ParsedCV:
    text: str
    facts: list[ExtractedFact] = field(default_factory=list)
    sections: dict[str, str] = field(default_factory=dict)


SECTION_ALIASES = {
    "summary": ["summary", "profile", "professional summary", "about me", "objective", "career objective"],
    "experience": [
        "experience", "work experience", "professional experience", "employment", "employment history",
        "work history", "career history", "internships", "internship experience",
    ],
    "education": ["education", "academic background", "academic qualifications", "qualifications"],
    "skills": ["skills", "technical skills", "core skills", "key skills", "competencies", "core competencies"],
    "certifications": ["certifications", "certificates", "licenses", "licenses & certifications", "courses"],
    "languages": ["languages", "language skills"],
    "projects": ["projects", "academic projects", "personal projects", "key projects"],
}
_HEADER_LOOKUP = {a: k for k, al in SECTION_ALIASES.items() for a in al}

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"(?<!\w)(\+?\d[\d\s().-]{7,}\d)(?!\w)")
LINKEDIN_RE = re.compile(r"(?:https?://)?(?:[a-z]{2,3}\.)?linkedin\.com/in/[A-Za-z0-9_%-]+/?", re.I)
GITHUB_RE = re.compile(r"(?:https?://)?github\.com/[A-Za-z0-9_-]+/?", re.I)
MONTHS = r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?"
DATE_TOKEN = rf"(?:{MONTHS}\s+\d{{4}}|\d{{1,2}}/\d{{4}}|\d{{4}})"
RANGE_RE = re.compile(
    rf"(?P<start>{DATE_TOKEN})\s*(?:-|–|—|to)\s*(?P<end>{DATE_TOKEN}|present|current|now|ongoing)", re.I
)
DEGREE_RE = re.compile(
    r"\b(b\.?sc|bsc|b\.?s\.?|b\.?a\.?|b\.?e\.?|b\.?tech|bachelor(?:'s)?|m\.?sc|msc|m\.?s\.?|m\.?a\.?|mba|m\.?tech|"
    r"master(?:'s)?|ph\.?d|doctorate|diploma|associate(?:'s)? degree|high school|a[- ]levels?)\b",
    re.I,
)
INSTITUTION_RE = re.compile(r"\b(university|college|institute|school|academy|polytechnic)\b", re.I)
YEAR_RE = re.compile(r"\b(19[6-9]\d|20[0-4]\d)\b")
_MONTH_NUM = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


def normalize_date(token: str) -> str | None:
    """'Mar 2021' -> '2021-03', '03/2021' -> '2021-03', '2021' -> '2021', present -> 'present'."""
    t = token.strip().lower().rstrip(".")
    if t in {"present", "current", "now", "ongoing"}:
        return "present"
    m = re.match(rf"({MONTHS})\s+(\d{{4}})", t)
    if m:
        return f"{m.group(2)}-{_MONTH_NUM[m.group(1)[:3]]:02d}"
    m = re.match(r"(\d{1,2})/(\d{4})", t)
    if m and 1 <= int(m.group(1)) <= 12:
        return f"{m.group(2)}-{int(m.group(1)):02d}"
    m = re.match(r"(\d{4})$", t)
    return m.group(1) if m else None


def _is_header(line: str) -> str | None:
    clean = re.sub(r"[^a-z& ]", "", line.strip().lower()).strip()
    if not clean or len(clean) > 40:
        return None
    return _HEADER_LOOKUP.get(clean)


def split_sections(text: str) -> dict[str, str]:
    sections: dict[str, list[str]] = {"_header": []}
    current = "_header"
    for line in text.splitlines():
        h = _is_header(line)
        if h:
            current = h
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(line)
    return {k: "\n".join(v).strip() for k, v in sections.items()}


def _split_list(block: str) -> list[str]:
    items: list[str] = []
    for line in block.splitlines():
        line = re.sub(r"^[\s•●▪◦*·\-–]+", "", line).strip()
        if not line:
            continue
        # "Category: a, b, c" -> items a, b, c
        if ":" in line and len(line.split(":", 1)[0]) < 40:
            line = line.split(":", 1)[1]
        for part in re.split(r"[,|;•·]", line):
            p = part.strip(" .")
            if 1 < len(p) <= 60:
                items.append(p)
    seen: set[str] = set()
    out = []
    for i in items:
        if i.lower() not in seen:
            seen.add(i.lower())
            out.append(i)
    return out


def _blocks(block: str) -> list[str]:
    """Split an experience/education section into entries at date-range lines or blank lines."""
    lines = [l.rstrip() for l in block.splitlines()]
    entries: list[list[str]] = []
    cur: list[str] = []
    for i, line in enumerate(lines):
        starts_new = bool(RANGE_RE.search(line)) and cur and any(RANGE_RE.search(x) for x in cur)
        if not line.strip():
            if cur:
                entries.append(cur)
                cur = []
            continue
        if starts_new:
            # the heading line(s) right before the date line belong to the new entry
            carry = []
            while cur and not RANGE_RE.search(cur[-1]) and len(cur) > 1 and not cur[-1].lstrip().startswith(("•", "-", "*", "●")):
                carry.insert(0, cur.pop())
                if len(carry) >= 2:
                    break
            entries.append(cur)
            cur = carry
        cur.append(line)
    if cur:
        entries.append(cur)
    return ["\n".join(e).strip() for e in entries if "\n".join(e).strip()]


def parse_cv_text(text: str) -> ParsedCV:
    parsed = ParsedCV(text=text, sections=split_sections(text))
    facts = parsed.facts
    header = parsed.sections.get("_header", "")
    head_lines = [l.strip() for l in header.splitlines() if l.strip()]

    # Name: first header line that looks like a personal name.
    for l in head_lines[:3]:
        if re.fullmatch(r"[A-Za-zÀ-ÿ'.-]+(?:\s+[A-Za-zÀ-ÿ'.-]+){1,4}", l) and not EMAIL_RE.search(l):
            facts.append(ExtractedFact("identity", "full_name", l, l))
            break

    if m := EMAIL_RE.search(text):
        facts.append(ExtractedFact("contact", "email", m.group(0), m.group(0)))
    for m in PHONE_RE.finditer(header or text[:1500]):
        digits = re.sub(r"\D", "", m.group(1))
        if 8 <= len(digits) <= 15 and not YEAR_RE.fullmatch(m.group(1).strip()):
            facts.append(ExtractedFact("contact", "phone", m.group(1).strip(), m.group(1)))
            break
    if m := LINKEDIN_RE.search(text):
        url = m.group(0)
        facts.append(ExtractedFact("contact", "linkedin_url", url if url.startswith("http") else "https://" + url, url))
    if m := GITHUB_RE.search(text):
        url = m.group(0)
        facts.append(ExtractedFact("contact", "github_url", url if url.startswith("http") else "https://" + url, url))

    if summary := parsed.sections.get("summary"):
        facts.append(ExtractedFact("narrative", "cv_summary", summary[:4000], summary[:300]))

    for skill in _split_list(parsed.sections.get("skills", "")):
        # Context (professional / academic / personal) cannot be known from a skills list.
        facts.append(ExtractedFact("skill", skill, "unknown", skill))
    for cert in _split_list(parsed.sections.get("certifications", "")):
        facts.append(ExtractedFact("certification", cert, "", cert))
    for lang in _split_list(parsed.sections.get("languages", "")):
        name, _, level = lang.partition("(")
        facts.append(ExtractedFact("language", name.strip(), level.strip(" )"), lang))

    for i, entry in enumerate(_blocks(parsed.sections.get("experience", "")), 1):
        g = f"emp:cv{i}"
        lines = [l.strip() for l in entry.splitlines() if l.strip()]
        rng = RANGE_RE.search(entry)
        heading = [l for l in lines[:3] if not l.startswith(("•", "-", "*", "●"))]
        heading_clean = [RANGE_RE.sub("", l).strip(" |,-–") for l in heading]
        heading_clean = [h for h in heading_clean if h]
        if heading_clean:
            parts = re.split(r"\s+(?:at|@)\s+|\s*[|,–—]\s*|\s+-\s+", heading_clean[0], maxsplit=1)
            facts.append(ExtractedFact("employment", "title", parts[0].strip(), heading[0], g))
            employer = parts[1].strip() if len(parts) > 1 else (heading_clean[1] if len(heading_clean) > 1 else "")
            if employer:
                facts.append(ExtractedFact("employment", "employer", employer, heading[0], g))
        if rng:
            s, e = normalize_date(rng.group("start")), normalize_date(rng.group("end"))
            if s:
                facts.append(ExtractedFact("employment", "start_date", s, rng.group(0), g))
            if e:
                facts.append(ExtractedFact("employment", "end_date", e, rng.group(0), g))
        bullets = [re.sub(r"^[•●▪*\-–]+\s*", "", l) for l in lines if l.startswith(("•", "-", "*", "●", "▪", "–"))]
        if bullets:
            facts.append(ExtractedFact("employment", "responsibilities", "\n".join(bullets), "\n".join(bullets)[:300], g))
        # Employment type (full-time / internship / academic ...) is never guessed.
        facts.append(ExtractedFact("employment", "raw_entry", entry, entry[:300], g))

    for i, entry in enumerate(_blocks(parsed.sections.get("education", "")), 1):
        g = f"edu:cv{i}"
        for line in entry.splitlines():
            line = line.strip()
            if not line:
                continue
            if DEGREE_RE.search(line) and not any(f.group_key == g and f.key == "degree" for f in facts):
                facts.append(ExtractedFact("education", "degree", RANGE_RE.sub("", line).strip(" |,-–"), line, g))
            if INSTITUTION_RE.search(line) and not any(f.group_key == g and f.key == "institution" for f in facts):
                inst = [p.strip() for p in re.split(r"[|,–—]| - ", line) if INSTITUTION_RE.search(p)]
                facts.append(ExtractedFact("education", "institution", inst[0] if inst else line, line, g))
        years = YEAR_RE.findall(entry)
        if years:
            facts.append(ExtractedFact("education", "graduation_year", max(years), entry[:200], g))
        facts.append(ExtractedFact("education", "raw_entry", entry, entry[:300], g))

    for i, entry in enumerate(_blocks(parsed.sections.get("projects", "")), 1):
        g = f"proj:cv{i}"
        first = entry.splitlines()[0].strip()
        facts.append(ExtractedFact("project", "name", RANGE_RE.sub("", first).strip(" |,-–"), first, g))
        facts.append(ExtractedFact("project", "description", entry, entry[:300], g))
    return parsed
